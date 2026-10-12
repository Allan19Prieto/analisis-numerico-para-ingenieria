import argparse
import hashlib
import json
import threading
from contextlib import contextmanager
from pathlib import Path
from time import perf_counter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import parte1
from parte1 import jacobi_valores_propios

ALTO = 56
ANCHO = 46
PERSONAS = 40
FOTOS_ENTRENAMIENTO = 9
ITER_MAX = 10**6
TOL = 1e-5
UMBRAL_RELATIVO = 1e-10
VARIACION_OBJETIVO = 0.95
VERSION_MODELO = "bloque3-directo-AAtranspuesta-v1"

JUSTIFICACION = (
    "Se seleccionó Jacobi cíclico a partir de la comparación de la Parte II. "
    "En la ejecución de referencia de la matriz de orden 200, Jacobi alcanzó "
    "la tolerancia en 11 barridos, con residuo máximo cercano a 2.86e-14. "
    "QR sin desplazamiento agotó 10000 pasos sin alcanzar la tolerancia "
    "y obtuvo un residuo máximo cercano a 8.59e-4. Jacobi también presentó "
    "menor tiempo en aquella ejecución. Estos resultados corresponden a "
    "otra matriz y motivan la elección, pero no garantizan la convergencia "
    "ni la exactitud de la clasificación actual. Por eso se comprueban "
    "también los pares propios calculados para G."
)


@contextmanager
def aviso_calculo():
    """Informa que el cálculo sigue activo, sin inventar un porcentaje."""
    detener = threading.Event()
    inicio = perf_counter()

    def informar():
        while not detener.wait(30):
            minutos = (perf_counter() - inicio) / 60
            print(f"  Jacobi sigue calculando: {minutos:.1f} min transcurridos.", flush=True)

    hilo = threading.Thread(target=informar, daemon=True)
    hilo.start()
    try:
        yield
    finally:
        detener.set()
        hilo.join()


def verificar_archivos(dataset):
    """Comprueba nombres, sin leer los píxeles de prueba para entrenar."""
    faltantes = [
        dataset / f"s{persona}" / f"{foto}.png"
        for persona in range(1, PERSONAS + 1)
        for foto in range(1, 11)
        if not (dataset / f"s{persona}" / f"{foto}.png").is_file()
    ]
    if faltantes:
        muestra = "\n".join(str(p) for p in faltantes[:10])
        raise FileNotFoundError(
            f"Faltan {len(faltantes)} imágenes. Primeras rutas esperadas:\n{muestra}\n"
            "Use --dataset para indicar la carpeta que contiene directamente s1 ... s40."
        )


def leer_vector(ruta):
    """Paso 3: intensidades [0,1], tamaño original y orden por columnas."""
    with Image.open(ruta) as foto:
        if foto.size != (ANCHO, ALTO):
            raise ValueError(
                f"{ruta}: tamaño {foto.size}; se esperaba ancho 46 y alto 56. "
                "No se redimensionan imágenes. Use la distribución de TEC Digital."
            )
        # Se permite RGB solamente si sus tres canales son idénticos.
        raw = np.asarray(foto)
        if foto.mode in ("RGB", "RGBA"):
            if not (np.array_equal(raw[..., 0], raw[..., 1]) and
                    np.array_equal(raw[..., 1], raw[..., 2])):
                raise ValueError(f"{ruta}: la fotografía no está en escala de grises.")
        elif foto.mode not in ("L", "P"):
            raise ValueError(f"{ruta}: se esperaba una imagen de grises de 8 bits.")
        if foto.mode == "P":
            rgb = np.asarray(foto.convert("RGB"))
            if not (np.array_equal(rgb[..., 0], rgb[..., 1]) and
                    np.array_equal(rgb[..., 1], rgb[..., 2])):
                raise ValueError(f"{ruta}: paleta con colores, no escala de grises.")
        imagen = np.asarray(foto.convert("L"), dtype=np.float64) / 255.0
    return imagen.reshape(-1, order="F")


def cargar_entrenamiento(dataset):
    """Orden numérico por persona y después por fotografía."""
    vectores, etiquetas, rutas = [], [], []
    for persona in range(1, PERSONAS + 1):
        for foto in range(1, FOTOS_ENTRENAMIENTO + 1):
            relativa = f"s{persona}/{foto}.png"
            vectores.append(leer_vector(dataset / relativa))
            etiquetas.append(persona)
            rutas.append(relativa)
    return np.column_stack(vectores), np.array(etiquetas), np.array(rutas)


def huella_entrenamiento(dataset):
    """La caché depende SOLO del entrenamiento, código y parámetros."""
    h = hashlib.sha256()
    h.update(VERSION_MODELO.encode())
    h.update(Path(parte1.__file__).read_bytes())
    h.update(Path(__file__).read_bytes())
    h.update(str((ALTO, ANCHO, ITER_MAX, TOL, UMBRAL_RELATIVO, VARIACION_OBJETIVO)).encode())
    for persona in range(1, PERSONAS + 1):
        for foto in range(1, FOTOS_ENTRENAMIENTO + 1):
            relativa = f"s{persona}/{foto}.png"
            h.update(relativa.encode())
            h.update((dataset / relativa).read_bytes())
    return h.hexdigest()


def entrenar(S, y, rutas, huella):
    """Pasos 4-7. Esta función no recibe fotografías ni etiquetas de prueba."""
    promedio = np.mean(S, axis=1)
    A = S - promedio[:, None]
    G = A @ A.T  # Exactamente la matriz indicada por la guía.
    print(f"Paso 5: G tiene tamaño {G.shape[0]} x {G.shape[1]}.", flush=True)
    print("La primera ejecución puede tardar mucho; no hay porcentaje disponible.", flush=True)

    inicio = perf_counter()
    with aviso_calculo():
        valores, U, erk, barridos, conv = jacobi_valores_propios(G, ITER_MAX, TOL)
    tiempo = perf_counter() - inicio

    if conv != 1 or not np.isfinite(erk):
        raise RuntimeError(
            f"Jacobi no alcanzó la tolerancia: erk={erk}, barridos={barridos}, conv={conv}. "
            "No se generarán predicciones como si el entrenamiento hubiera convergido."
        )
    if not np.all(np.isfinite(valores)) or not np.all(np.isfinite(U)):
        raise RuntimeError("El cálculo produjo valores no finitos.")
    if valores[0] <= 0:
        raise ValueError("No hay variación positiva en el entrenamiento para construir caras base.")

    conservar = valores > UMBRAL_RELATIVO * valores[0]
    valores_r = valores[conservar]
    Ur = U[:, conservar]
    r = len(valores_r)
    if r == 0:
        raise ValueError("No quedaron valores propios después del filtrado.")

    # Paso 6: menor k que conserva al menos el 95% de la variación retenida.
    acumulada = np.cumsum(valores_r) / np.sum(valores_r)
    k = int(np.searchsorted(acumulada, VARIACION_OBJETIVO, side="left") + 1)
    Uk = Ur[:, :k].copy()

    # Comprobaciones de TODOS los vectores retenidos, no solo los primeros seis.
    residuos = np.linalg.norm(G @ Ur - Ur * valores_r[None, :], axis=0)
    erort = float(np.linalg.norm(Ur.T @ Ur - np.eye(r), "fro"))
    error_norma = float(np.max(np.abs(np.linalg.norm(Ur, axis=0) - 1.0)))
    # G es semidefinida positiva: usamos lambda1 como escala del residuo.
    residuo_relativo = float(np.max(residuos) / valores[0])
    avisos = []
    if r > S.shape[1] - 1:
        avisos.append(
            f"Se retuvieron {r} valores, aunque centrar 360 columnas limita el rango exacto "
            "a 359. Revise la precisión de los pares propios cerca del umbral."
        )
    if residuo_relativo > 1e-6:
        avisos.append(
            f"El residuo máximo relativo es {residuo_relativo:.3e}; el cambio de valores "
            "propios no garantiza por sí solo precisión suficiente."
        )
    if erort > 1e-6:
        avisos.append(f"El error de ortogonalidad de las caras retenidas es {erort:.3e}.")
    if np.min(valores) < -UMBRAL_RELATIVO * valores[0]:
        avisos.append("Se observaron valores negativos mayores que el umbral de redondeo esperado.")

    # Paso 7: las columnas de X conservan las etiquetas y las rutas de S.
    X = Uk.T @ A
    return {
        "promedio": promedio, "Uk": Uk, "X": X, "y": y, "rutas": rutas,
        "valores": valores_r, "acumulada": acumulada,
        "primeras_caras": Ur[:, :min(6, r)].copy(),
        "r": np.array(r), "k": np.array(k), "proporcion": np.array(acumulada[k-1]),
        "erk": np.array(erk), "barridos": np.array(barridos), "conv": np.array(conv),
        "tiempo": np.array(tiempo), "residuos": residuos,
        "erort": np.array(erort), "error_norma": np.array(error_norma),
        "residuo_relativo": np.array(residuo_relativo),
        "avisos": np.array(avisos, dtype=str), "huella": np.array(huella),
    }


def obtener_modelo(dataset, salida, S, y, rutas, reentrenar):
    huella = huella_entrenamiento(dataset)
    archivo = salida / "modelo.npz"
    if archivo.is_file() and not reentrenar:
        try:
            with np.load(archivo, allow_pickle=False) as guardado:
                if str(guardado["huella"].item()) == huella:
                    modelo = {clave: guardado[clave].copy() for clave in guardado.files}
                    print("Modelo del mismo entrenamiento recuperado de modelo.npz.", flush=True)
                    return modelo, True
            print("Cambió el entrenamiento o el código: se recalculará el modelo.", flush=True)
        except (ValueError, OSError, KeyError) as exc:
            print(f"No se pudo reutilizar la caché: {exc}. Se recalculará.", flush=True)
    modelo = entrenar(S, y, rutas, huella)
    # Escritura temporal: un corte al guardar no destruye la caché anterior.
    temporal = salida / "modelo_temporal.npz"
    np.savez_compressed(temporal, **modelo)
    temporal.replace(archivo)
    return modelo, False


def evaluar_pruebas(dataset, modelo):
    """Paso 8: nearest neighbor; la etiqueta real NO participa en predecir."""
    resultados, pruebas = [], []
    for persona in range(1, PERSONAS + 1):
        archivo = f"s{persona}/10.png"
        f = leer_vector(dataset / archivo)
        x = modelo["Uk"].T @ (f - modelo["promedio"])
        distancias = np.linalg.norm(modelo["X"] - x[:, None], axis=0)
        indice = int(np.argmin(distancias))  # Primer mínimo en caso de empate.
        predicha = int(modelo["y"][indice])
        # Solo después de predecir se utiliza la etiqueta real para evaluar.
        correctos = np.flatnonzero(modelo["y"] == persona)
        indice_real = int(correctos[np.argmin(distancias[correctos])])
        resultados.append({
            "fotografia": archivo, "real": persona, "predicha": predicha,
            "correcta": predicha == persona, "indice": indice,
            "cercana": str(modelo["rutas"][indice]), "distancia": float(distancias[indice]),
            "cercana_real": str(modelo["rutas"][indice_real]),
            "distancia_real": float(distancias[indice_real]),
        })
        pruebas.append(f)
    return resultados, np.column_stack(pruebas)


def como_imagen(vector):
    return vector.reshape(ALTO, ANCHO, order="F")


def crear_figuras(modelo, S, pruebas, resultados, salida):
    carpeta = salida / "figuras"
    carpeta.mkdir(exist_ok=True)
    caras = modelo["primeras_caras"]
    fig, ejes = plt.subplots(2, 4, figsize=(10.0, 6.4))
    for ax in ejes.flat:
        ax.axis("off")
    ejes.flat[0].imshow(como_imagen(modelo["promedio"]), cmap="gray", vmin=0, vmax=1)
    ejes.flat[0].set_title("Rostro promedio")
    for j in range(caras.shape[1]):
        ax = ejes.flat[j+1]
        limite = float(np.max(np.abs(caras[:, j])))
        ax.imshow(como_imagen(caras[:, j]), cmap="gray", vmin=-limite, vmax=limite)
        ax.set_title(f"Cara base {j+1}")
    fig.suptitle("Promedio y primeras caras base del entrenamiento", fontsize=14)
    fig.text(.5, .015, "Caras base con contraste propio; el signo puede invertirse sin cambiar su validez.",
             ha="center", fontsize=9)
    fig.tight_layout(rect=(0, .04, 1, .94), h_pad=2.5)
    fig.savefig(carpeta / "01_promedio_y_caras.png", dpi=170)
    plt.close(fig)

    k = int(modelo["k"])
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.plot(np.arange(1, len(modelo["acumulada"])+1), modelo["acumulada"]*100, color="#245D83")
    ax.axhline(95, color="#9F422D", linestyle="--", label="95% requerido")
    ax.axvline(k, color="#28765D", linestyle=":", label=f"k = {k}")
    ax.scatter([k], [float(modelo["proporcion"])*100], color="#28765D", zorder=5)
    ax.set(xlabel="Número de caras base", ylabel="Variación acumulada (%)",
           title="Selección de k usando solo entrenamiento", ylim=(0, 101))
    ax.grid(alpha=.2); ax.legend(); fig.tight_layout()
    fig.savefig(carpeta / "02_variacion.png", dpi=170)
    plt.close(fig)

    paginas = []
    for inicio in range(0, len(resultados), 10):
        lote = resultados[inicio:inicio+10]
        fig, ejes = plt.subplots(5, 4, figsize=(10, 12))
        for ax in ejes.flat:
            ax.axis("off")
        for local, fila in enumerate(lote):
            renglon, grupo = divmod(local, 2)
            ax_prueba, ax_cercana = ejes[renglon, 2*grupo:2*grupo+2]
            i = inicio + local
            ax_prueba.imshow(como_imagen(pruebas[:, i]), cmap="gray", vmin=0, vmax=1)
            ax_prueba.set_title(f'Prueba s{fila["real"]}/10.png', fontsize=9)
            ax_cercana.imshow(como_imagen(S[:, fila["indice"]]), cmap="gray", vmin=0, vmax=1)
            estado = "Correcta" if fila["correcta"] else "Incorrecta"
            ax_cercana.set_title(f'{fila["cercana"]}\n{estado} - predicción {fila["predicha"]}',
                                fontsize=9, color="#236647" if fila["correcta"] else "#A33726")
        fig.suptitle(f'Fotografías de prueba y entrenamiento más cercano\nPersonas {inicio+1} a {inicio+len(lote)}',
                     fontsize=13)
        fig.tight_layout(rect=(0, 0, 1, .94))
        nombre = carpeta / f"comparaciones_{inicio+1:02d}_{inicio+len(lote):02d}.png"
        fig.savefig(nombre, dpi=160); plt.close(fig)
        paginas.append(nombre)
    return carpeta / "01_promedio_y_caras.png", carpeta / "02_variacion.png", paginas


def analizar(modelo, resultados):
    aciertos = sum(f["correcta"] for f in resultados)
    exactitud = 100 * aciertos / len(resultados)
    r, k = int(modelo["r"]), int(modelo["k"])
    textos = [
        f"Se identificaron correctamente {aciertos} de {len(resultados)} fotografías de prueba. "
        f"La exactitud fue {exactitud:.2f}%. Cada fotografía correspondió a una persona conocida "
        "a través de otras nueve imágenes de entrenamiento.",
        f"El filtrado lambda_j > 1e-10 * lambda_1 retuvo r = {r} valores propios. "
        f"El menor número que conservó al menos el 95% de la variación retenida fue k = {k}. "
        f"La proporción obtenida fue {float(modelo['proporcion'])*100:.6f}%. "
        "k se fijó antes de leer las fotografías de prueba y no se ajustó según los aciertos.",
        f"Las fotografías pasaron de 2576 intensidades a {k} coordenadas. Los vectores propios "
        "definen las direcciones de representación y sus valores propios cuantifican la variación "
        "asociada. Retener el 95% de esa variación no equivale a asegurar 95% de exactitud: "
        "una dirección con mucha variación no necesariamente separa identidades.",
        f"Jacobi terminó en {int(modelo['barridos'])} barridos con erk = {float(modelo['erk']):.6e}. "
        f"En las caras retenidas, el residuo máximo fue {np.max(modelo['residuos']):.6e}, "
        f"su valor relativo a lambda_1 fue {float(modelo['residuo_relativo']):.6e} y el error "
        f"de ortogonalidad fue {float(modelo['erort']):.6e}. La convergencia del cambio de "
        "valores propios y la precisión de los pares propios son comprobaciones diferentes.",
    ]
    errores = [f for f in resultados if not f["correcta"]]
    if errores:
        textos.append(
            f"Se observaron {len(errores)} confusiones. En cada caso la fotografía de otra "
            "identidad obtuvo una distancia menor o igual que las fotos de la persona real. "
            "Las comparaciones siguientes permiten revisarlas. La iluminación, la expresión o "
            "la pose podrían influir, pero el programa no determina automáticamente cuál fue "
            "la causa visual de cada error."
        )
        for f in errores:
            textos.append(
                f'{f["fotografia"]}: persona real {f["real"]}, predicción {f["predicha"]}; '
                f'la imagen elegida fue {f["cercana"]}, a distancia {f["distancia"]:.6f}. '
                f'La imagen más cercana de la identidad real fue {f["cercana_real"]}, '
                f'a distancia {f["distancia_real"]:.6f}. '
                'La comparación se realizó en el espacio de caras base, no entre píxeles directamente.'
            )
    else:
        textos.append(
            "No se observaron identificaciones incorrectas en estas 40 pruebas. Esto no "
            "garantiza acierto en fotografías nuevas ni en condiciones de captura diferentes."
        )
    textos.extend(str(a) for a in modelo["avisos"])
    textos.append(
        "El procedimiento asigna siempre una de las 40 identidades conocidas. No incluye "
        "detección de rostros, rechazo de personas desconocidas ni ajuste de k con datos de prueba. "
        "Los resultados describen únicamente esta separación fija de entrenamiento y prueba."
    )
    return aciertos, exactitud, textos


def guardar_texto(modelo, resultados, textos, aciertos, exactitud, salida, cache):
    encabezado = (
        "PARTE III - RECONOCIMIENTO FACIAL\n\n" + JUSTIFICACION + "\n\n"
        f"Entrenamiento: 360 imágenes. Prueba: 40 imágenes. Resolución: 56 x 46.\n"
        f"Método: Jacobi cíclico. iterMax={ITER_MAX}; tol={TOL}.\n"
        f"r={int(modelo['r'])}; k={int(modelo['k'])}; "
        f"variación={float(modelo['proporcion'])*100:.6f}%.\n"
        f"erk={float(modelo['erk']):.8e}; barridos={int(modelo['barridos'])}; "
        f"conv={int(modelo['conv'])}.\n"
        f"Tiempo del cálculo original de Jacobi: {float(modelo['tiempo']):.3f} s.\n"
        f"Modelo reutilizado desde caché: {'sí' if cache else 'no'}.\n"
        f"Aciertos={aciertos}/40; exactitud={exactitud:.2f}%.\n\n"
    )
    filas = [f'{"Fotografía":<16} {"Real":>5} {"Predicha":>9} {"Correcta":>9} {"Más cercana":>17} {"Distancia":>13}']
    for f in resultados:
        filas.append(
            f'{f["fotografia"]:<16} {f["real"]:5d} {f["predicha"]:9d} '
            f'{("Sí" if f["correcta"] else "No"):>9} {f["cercana"]:>17} {f["distancia"]:13.6f}'
        )
    informe = encabezado + "\n".join(filas) + "\n\nANÁLISIS\n\n" + "\n\n".join(textos)
    (salida / "resultados.txt").write_text(informe + "\n", encoding="utf-8")
    print(informe)
    metricas = {
        "metodo": "Jacobi cíclico directo sobre A @ A.T", "iterMax": ITER_MAX, "tol": TOL,
        "r": int(modelo["r"]), "k": int(modelo["k"]),
        "proporcion": float(modelo["proporcion"]), "aciertos": aciertos,
        "exactitud_porcentaje": exactitud, "erk": float(modelo["erk"]),
        "barridos": int(modelo["barridos"]), "conv": int(modelo["conv"]),
        "residuo_maximo": float(np.max(modelo["residuos"])),
        "residuo_relativo": float(modelo["residuo_relativo"]),
        "error_ortogonalidad": float(modelo["erort"]),
        "tiempo_jacobi_segundos": float(modelo["tiempo"]),
        "modelo_reutilizado": cache, "huella_entrenamiento": str(modelo["huella"].item()),
        "advertencias_numericas": modelo["avisos"].tolist(), "predicciones": resultados,
    }
    (salida / "metricas.json").write_text(json.dumps(metricas, ensure_ascii=False, indent=2), encoding="utf-8")


def generar_word(modelo, resultados, textos, aciertos, exactitud, figuras, salida, integrantes, grupo, cache):
    """Genera un Word editable con tablas nativas y las figuras de esta ejecución."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.7)
    for nombre in ("Normal", "Title", "Heading 1", "Heading 2"):
        estilo = doc.styles[nombre]
        estilo.font.name = "Calibri"
        estilo.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.core_properties.title = "Reconocimiento facial mediante valores propios"
    doc.core_properties.author = integrantes or "Por completar"

    def p(texto):
        return doc.add_paragraph(texto)

    def h(texto):
        return doc.add_heading(texto, level=1)

    def imagen(ruta, ancho=7.0, alto=6.9):
        with Image.open(ruta) as im:
            w, ht = im.size
        escala = min(ancho / w, alto / ht)
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(ruta), width=Inches(w*escala), height=Inches(ht*escala))

    def tabla(datos, anchos):
        t = doc.add_table(rows=1, cols=len(anchos))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, ancho in zip(t.columns, anchos):
            col.width = Inches(ancho)
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:"+lado)
            for clave, valor in (("val","single"),("sz","4"),("color","D9D9D9")):
                el.set(qn("w:"+clave), valor)
            bordes.append(el)
        t._tbl.tblPr.append(bordes)
        for i, fila in enumerate(datos):
            celdas = t.rows[0].cells if i == 0 else t.add_row().cells
            trpr = t.rows[i]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j, (celda, valor) in enumerate(zip(celdas, fila)):
                celda.width = Inches(anchos[j])
                celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                celda.text = str(valor)
                props = celda._tc.get_or_add_tcPr()
                fondo = OxmlElement("w:shd")
                fondo.set(qn("w:fill"), "243E51" if i == 0 else ("F0F3F5" if i % 2 else "FFFFFF"))
                props.append(fondo)
                marg = OxmlElement("w:tcMar")
                for lado in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:"+lado)
                    el.set(qn("w:w"), "70")
                    el.set(qn("w:type"), "dxa")
                    marg.append(el)
                props.append(marg)
                for par in celda.paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.line_spacing = 1
                    if j > 0:
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in par.runs:
                        run.font.size = Pt(9)
                        run.bold = i == 0
                        run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
        return t

    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pie.add_run("Bloque 3 · Reconocimiento facial · Página ").font.size = Pt(8)
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)
    doc.add_paragraph("Reconocimiento facial mediante valores propios", style="Title")
    p("CE-1111 · Portafolio computacional · Bloque 3 · Parte III")
    p(f"Integrantes: {integrantes or 'Por completar'} | Grupo: {grupo or 'Por completar'}")
    p(f"Se evalúa la identificación de 40 personas mediante caras base. Se obtuvieron "
      f"{aciertos} aciertos en 40 pruebas, para una exactitud de {exactitud:.2f}%. "
      "Este informe documenta el entrenamiento, las predicciones y sus limitaciones.")
    h("Método numérico seleccionado")
    p(JUSTIFICACION)
    h("Resultados del entrenamiento y la evaluación")
    datos = [["Medida", "Resultado"],
             ["Imágenes de entrenamiento / prueba", "360 / 40"],
             ["Dimensión de G = A @ A.T", "2576 x 2576"],
             ["iterMax / tolerancia", "1000000 / 1e-5"],
             ["Valores retenidos r / caras seleccionadas k", f"{int(modelo['r'])} / {int(modelo['k'])}"],
             ["Variación conservada", f"{float(modelo['proporcion'])*100:.6f}%"],
             ["Aciertos / exactitud", f"{aciertos}/40 / {exactitud:.2f}%"],
             ["Barridos / conv", f"{int(modelo['barridos'])} / {int(modelo['conv'])}"],
             ["Error final erk", f"{float(modelo['erk']):.6e}"],
             ["Residuo máximo retenido", f"{np.max(modelo['residuos']):.6e}"],
             ["Error de ortogonalidad retenido", f"{float(modelo['erort']):.6e}"],
             ["Tiempo original de Jacobi", f"{float(modelo['tiempo']):.3f} s"],
             ["Modelo recuperado de caché", "Sí" if cache else "No"]]
    tabla(datos, [4.5, 2.6])
    doc.add_page_break()
    h("Rostro promedio y caras base")
    promedio, variacion, comparaciones = figuras
    imagen(promedio, alto=3.7)
    p("Se muestran las primeras seis caras disponibles; el signo de un vector propio puede invertirse sin cambiar su validez.")
    h("Selección del número de caras base")
    imagen(variacion, alto=2.6)
    p("k es el menor entero que conserva al menos el 95% de la variación retenida. "
      "El promedio, las caras base y k se calcularon solo con las fotografías 1 a 9. "
      "Las fotografías 10 se leyeron después para evaluar el modelo.")
    doc.add_page_break()
    h("Predicciones de las 40 personas")
    p("La predicción es la etiqueta de la fotografía de entrenamiento de menor distancia "
      "euclidiana en el espacio de caras base. Un empate exacto se resuelve con la primera "
      "imagen en orden numérico de persona y fotografía.")
    datos = [["Fotografía", "Real", "Predicha", "Correcta", "Entrenamiento cercano"]]
    for f in resultados:
        datos.append([f["fotografia"], str(f["real"]), str(f["predicha"]),
                      "Sí" if f["correcta"] else "No", f["cercana"]])
    tabla(datos, [1.5, .6, .9, .9, 3.2])
    for i, ruta in enumerate(comparaciones):
        doc.add_page_break()
        h(f"Comparación visual de las personas {i*10+1} a {i*10+10}")
        imagen(ruta, alto=8.3)
    doc.add_page_break()
    h("Análisis e interpretación")
    for texto in textos:
        p(texto)
    h("Fuentes y procedimiento")
    p("Soto Quirós, J. P. Portafolio Computacional, CE-1111, II semestre 2026, "
      "Bloque 3, Parte III. Enunciado suministrado por el curso.")
    p("Reconocimiento facial mediante valores y vectores propios. Procedimiento "
      "con la base de datos AT&T. Guía experimental suministrada por el curso, pasos 1 al 8.")
    p("Fotografías: distribución dataset_att.zip proporcionada en TEC Digital, "
      "basada en AT&T Database of Faces. Resolución de 56 filas por 46 columnas.")
    # Eliminar bordes decorativos heredados de la plantilla de Word.
    for raiz in (doc.styles.element, doc.element):
        for borde in list(raiz.iter(qn("w:pBdr"))):
            borde.getparent().remove(borde)
    doc.save(str(salida / "resultados.docx"))


def main():
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", type=Path, default=base / "dataset",
                        help="Carpeta que contiene directamente s1 ... s40.")
    parser.add_argument("--salida", type=Path, default=base / "resultados_parte3")
    parser.add_argument("--reentrenar", action="store_true", help="Ignorar la caché y recalcular G y Jacobi.")
    parser.add_argument("--integrantes", default="", help="Nombres para el informe Word.")
    parser.add_argument("--grupo", default="", help="Número de grupo para el informe Word.")
    args = parser.parse_args()
    dataset, salida = args.dataset.resolve(), args.salida.resolve()
    verificar_archivos(dataset)
    salida.mkdir(parents=True,exist_ok=True)
    print("Pasos 1-3: leyendo las 360 imágenes de entrenamiento...", flush=True)
    S,y,rutas = cargar_entrenamiento(dataset)
    print("Pasos 4-7: promedio, caras base y coordenadas de entrenamiento...", flush=True)
    modelo, cache = obtener_modelo(dataset,salida,S,y,rutas,args.reentrenar)
    print(f"Modelo fijado: r={int(modelo['r'])}, k={int(modelo['k'])}. Ahora se leerán las pruebas.", flush=True)
    resultados,pruebas = evaluar_pruebas(dataset,modelo)
    aciertos,exactitud,textos = analizar(modelo,resultados)
    figuras = crear_figuras(modelo,S,pruebas,resultados,salida)
    guardar_texto(modelo,resultados,textos,aciertos,exactitud,salida,cache)
    generar_word(modelo,resultados,textos,aciertos,exactitud,figuras,salida,
                args.integrantes,args.grupo,cache)
    print(f"\nInforme final: {salida / 'resultados.docx'}")
    print(f"Figuras, modelo y resultados de texto: {salida}")


if __name__ == "__main__":
    try:
        main()
    except (FileNotFoundError, ValueError, RuntimeError, OSError) as exc:
        raise SystemExit(f"\nNo se completó la parte III: {exc}") from exc
